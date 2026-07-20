"""Shared logic for member bulk upload — used by both the general
(/admin/members/bulk-upload, multi-partner) and partner-scoped
(/admin/partners/{id}/members/bulk-upload) endpoints, so the two never
drift out of sync on columns/validation again.
"""
import re
import io
import datetime as _dt
from typing import Dict, List, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from fastapi import HTTPException

from ..db.queries.partner_query import PartnerQuery
from ..db.queries.plan_query import PlanQuery

MEMBER_COL_MAP = {
    # name variants
    "primary member full name": "name", "full name": "name", "name": "name",
    "member name": "name",
    # email variants
    "primary email id": "email", "email id": "email", "email": "email",
    # mobile variants
    "primary mobile no.": "mobile_no", "primary mobile no": "mobile_no",
    "mobile number": "mobile_no", "mobile no": "mobile_no",
    "mobile no.": "mobile_no", "mobile": "mobile_no",
    "phone": "mobile_no", "phone number": "mobile_no",
    # gender
    "gender": "gender",
    # address variants
    "address line1": "address_line", "address line 1": "address_line",
    "address": "address_line", "address line": "address_line",
    # city / state / pin
    "city": "address_city",
    "state": "address_state",
    "pin code": "address_pin", "pin": "address_pin",
    "pincode": "address_pin", "postal code": "address_pin",
    # optional fields
    "sale date": "sale_date",
    "sales channel": "sales_channel",
    "partner branch code": "branch_code", "branch code": "branch_code",
    "sales person name": "salesperson_name", "salesperson name": "salesperson_name",
    "salesperson": "salesperson_name",
    "employee code": "employee_code",
    "data 1": "data1", "data1": "data1",
    "data 2": "data2", "data2": "data2",
    "data 3": "data3", "data3": "data3",
    # plan
    "plan": "plan_name", "plan name": "plan_name",
    "plan code": "plan_code", "plancode": "plan_code",
    # partner
    "partner code": "partner_code", "partnercode": "partner_code",
}

MEMBER_MANDATORY = {"name", "mobile_no", "email", "plan_code"}

_MOB_RE = re.compile(r"^\+?[\d\s\-()]{7,15}$")
_PIN_RE = re.compile(r"^\d{6}$")


def resolve_partner_for_row(partner_code: Optional[str], fallback_partner_id: Optional[str]) -> Tuple[Optional[str], Optional[str]]:
    """Resolve the partner to enroll this row under via its 'Partner Code' column,
    falling back to a single pre-selected partner_id when the file has no per-row
    partner code. Either way, the resolved partner must be Active.
    Returns (partner_id, error_message)."""
    if partner_code:
        partner = PartnerQuery().get_by_code(partner_code)
        if not partner:
            return None, f"Partner Code '{partner_code}' not found"
        if partner.status != "Active":
            article = "an" if partner.status.lower().startswith(("i", "a", "e", "o", "u")) else "a"
            return None, f"Partner Code '{partner_code}' refers to {article} {partner.status.lower()} partner"
        return str(partner.id), None
    if fallback_partner_id:
        partner = PartnerQuery().get_by_id(fallback_partner_id)
        if not partner:
            return None, "Selected partner not found"
        if partner.status != "Active":
            return None, f"Selected partner is {partner.status.lower()}"
        return str(partner.id), None
    return None, "Partner Code is required"


def resolve_plan_for_row(partner_id: str, plan_code: Optional[str]) -> Tuple[Optional[str], Optional[str]]:
    """Resolve the plan to enroll this row into via its mandatory 'Plan Code' column.
    Returns (plan_id, error_message)."""
    from ..db.models.partner import PartnerPlan
    from ..db.session import session_scope

    if not plan_code:
        return None, "Plan Code is required"
    plan = PlanQuery().get_by_code(plan_code)
    if not plan or plan.status != "Active":
        return None, f"Plan Code '{plan_code}' not found or not active"
    if plan.plan_type == "partner":
        with session_scope() as s:
            linked = s.query(PartnerPlan).filter(
                PartnerPlan.partner_id == partner_id, PartnerPlan.plan_id == plan.id,
            ).first()
        if not linked:
            return None, f"Plan Code '{plan_code}' is not assigned to this partner"
    return str(plan.id), None


def parse_member_excel(contents: bytes, require_partner_code: bool) -> tuple:
    """Returns (header_row, data_rows, col_idx). Raises HTTPException on bad
    file or missing mandatory columns."""
    try:
        wb = openpyxl.load_workbook(filename=io.BytesIO(contents), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=422, detail="Invalid Excel file — must be .xlsx format")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise HTTPException(status_code=422, detail="Excel file is empty")

    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    col_idx = {}
    for i, h in enumerate(header):
        mapped = MEMBER_COL_MAP.get(h)
        if mapped:
            col_idx[mapped] = i

    missing = MEMBER_MANDATORY - set(col_idx.keys())
    if missing:
        raise HTTPException(status_code=422, detail=f"Missing mandatory columns: {sorted(missing)}")
    if require_partner_code and "partner_code" not in col_idx:
        raise HTTPException(
            status_code=422,
            detail="Missing mandatory column: ['partner_code'] (or select a partner before uploading)",
        )
    return rows[0], rows[1:], col_idx


def member_row_to_dict(row, col_idx: dict) -> dict:
    def cell(field):
        idx = col_idx.get(field)
        if idx is None:
            return None
        v = row[idx] if idx < len(row) else None
        return str(v).strip() if v is not None else None

    return {f: cell(f) for f in [
        "name", "email", "mobile_no", "gender", "address_line", "address_city",
        "address_state", "address_pin", "sale_date", "sales_channel", "branch_code",
        "salesperson_name", "employee_code", "data1", "data2", "data3",
        "plan_code", "partner_code",
    ]}


def validate_member_row(rd: dict) -> List[str]:
    errors = []
    for f in ("name", "mobile_no", "email", "plan_code"):
        if not rd.get(f):
            errors.append(f"{f.replace('_', ' ').title()} is required")
    mob = rd.get("mobile_no") or ""
    if mob and not _MOB_RE.match(mob):
        errors.append("Invalid Mobile Number format")
    pin = rd.get("address_pin") or ""
    if pin and not _PIN_RE.match(str(pin)):
        errors.append("Invalid PIN Code — must be 6 digits")
    gender = rd.get("gender") or ""
    if gender and gender not in ("Male", "Female", "Other"):
        errors.append("Gender must be Male, Female, or Other")
    return errors


def process_member_bulk_upload(contents: bytes, fallback_partner_id: Optional[str]) -> dict:
    """Parses the workbook and creates each valid row's member + enrollment.
    Returns {total_rows, created, skipped, errors, created_by_partner}."""
    from .member_service import MemberService
    from ..schemas.member import MemberCreate

    header_row, data_rows, col_idx = parse_member_excel(contents, require_partner_code=not fallback_partner_id)

    svc = MemberService()
    results = {"total_rows": len(data_rows), "created": [], "skipped": [], "errors": []}
    created_by_partner: Dict[str, int] = {}

    for row_num, row in enumerate(data_rows, start=2):
        rd = member_row_to_dict(row, col_idx)
        email = rd.get("email")
        if not email:
            results["skipped"].append({"row": row_num, "reason": "empty email"})
            continue

        row_errors = validate_member_row(rd)
        if row_errors:
            results["errors"].append({"row": row_num, "email": email, "reason": "; ".join(row_errors)})
            continue

        resolved_partner_id, partner_error = resolve_partner_for_row(rd.get("partner_code"), fallback_partner_id)
        if partner_error:
            results["errors"].append({"row": row_num, "email": email, "reason": partner_error})
            continue

        resolved_plan_id, plan_error = resolve_plan_for_row(resolved_partner_id, rd.get("plan_code"))
        if plan_error:
            results["errors"].append({"row": row_num, "email": email, "reason": plan_error})
            continue

        sale_date = None
        if rd.get("sale_date"):
            try:
                sale_date = _dt.date.fromisoformat(rd["sale_date"])
            except Exception:
                pass

        try:
            member_data = MemberCreate(
                email=email, name=rd.get("name"), mobile_no=rd.get("mobile_no"), gender=rd.get("gender"),
                address_line=rd.get("address_line"), address_city=rd.get("address_city"),
                address_state=rd.get("address_state"), address_pin=rd.get("address_pin"),
                partner_id=resolved_partner_id, plan_id=resolved_plan_id,
                sale_date=sale_date,
                sales_channel=rd.get("sales_channel"),
                branch_code=rd.get("branch_code"),
                salesperson_name=rd.get("salesperson_name"),
                employee_code=rd.get("employee_code"),
                data1=rd.get("data1"), data2=rd.get("data2"), data3=rd.get("data3"),
            )
            result = svc.create_member(member_data)
            results["created"].append({"row": row_num, "email": email, "id": str(result["user"].id)})
            created_by_partner[resolved_partner_id] = created_by_partner.get(resolved_partner_id, 0) + 1
        except HTTPException as e:
            results["skipped"].append({"row": row_num, "email": email, "reason": e.detail})
        except Exception as e:
            results["errors"].append({"row": row_num, "email": email, "reason": str(e)})

    results["created_by_partner"] = created_by_partner
    return results


def generate_sample_workbook(partners: list) -> io.BytesIO:
    """Builds the downloadable sample .xlsx for one or more (already Active-filtered)
    partners — a Members sheet with a sample row, and a Partner Plans reference sheet
    listing each partner's name/code alongside their active plans' name/code."""
    plan_q = PlanQuery()
    plan_rows = []  # (partner_name, partner_code, plan_name, plan_code, plan_status)
    for partner in partners:
        for plan in plan_q.list_for_partner(str(partner.id), active_only=True):
            plan_rows.append((partner.name, partner.partner_code, plan.name, plan.plan_code, plan.status))

    sample_partner_code = partners[0].partner_code if partners else "PTR001"
    sample_plan_code = plan_rows[0][3] if plan_rows else "PLN001"

    HEADERS = [
        "Email ID", "Name", "Mobile Number", "Gender", "Partner Code", "Plan Code",
        "Address", "City", "State", "PIN Code",
        "Sale Date", "Sales Channel", "Branch Code", "Salesperson Name", "Employee Code",
        "Data 1", "Data 2", "Data 3",
    ]
    SAMPLE = [
        "john.doe@example.com", "John Doe", "9876543210", "Male", sample_partner_code, sample_plan_code,
        "123 MG Road", "Mumbai", "Maharashtra", "400001",
        "2024-01-15", "Direct", "BRN001", "Jane Smith", "EMP123",
        "", "", "",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Members"
    hdr_fill = PatternFill(start_color="0A2257", end_color="0A2257", fill_type="solid")
    for col_num, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_num, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(len(h) + 4, 18)
    for col_num, v in enumerate(SAMPLE, start=1):
        ws.cell(row=2, column=col_num, value=v)

    ws2 = wb.create_sheet("Partner Plans")
    plan_headers = ["Partner Name", "Partner Code", "Plan Name", "Plan Code", "Status"]
    for col_num, h in enumerate(plan_headers, start=1):
        cell = ws2.cell(row=1, column=col_num, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
        ws2.column_dimensions[cell.column_letter].width = max(len(h) + 4, 18)
    if plan_rows:
        for row_num, (p_name, p_code, pl_name, pl_code, pl_status) in enumerate(plan_rows, start=2):
            ws2.cell(row=row_num, column=1, value=p_name)
            ws2.cell(row=row_num, column=2, value=p_code)
            ws2.cell(row=row_num, column=3, value=pl_name)
            ws2.cell(row=row_num, column=4, value=pl_code)
            ws2.cell(row=row_num, column=5, value=pl_status)
    else:
        ws2.cell(row=2, column=1, value="No active plans are assigned to the selected partner(s) yet.")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
