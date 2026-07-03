import io
import re
from uuid import UUID
import openpyxl
from openpyxl.styles import PatternFill, Font
from fastapi import APIRouter, Depends, File, Form, Request, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from typing import Optional
from ...schemas.base import ResponseModel
from ...schemas.partner import PartnerCreate, PartnerUpdate
from ...schemas.list_request import PartnerListRequest, MemberListRequest, PolicyListRequest
from ...services.partner_service import PartnerService
from ...services.plan_service import PlanService
from ...services.audit_service import AuditService
from ...db.queries.partner_query import PartnerQuery
from ...db.queries.user_query import UserQuery
from ...db.queries.member_query import MemberQuery
from ...db.queries.policy_query import PolicyQuery
from ...db.queries.policy_type_query import PolicyTypeQuery
from ..users import _require_superadmin
from ..plans import _plan_to_dict
from fastapi import HTTPException

admin_partners_router = APIRouter()


def _partner_dict(p, user=None) -> dict:
    return {
        "id": str(p.id),
        "user_id": str(p.user_id),
        "name": p.name,
        "partner_type": p.partner_type,
        "city": p.city,
        "state": getattr(p, "state", None),
        "legal_company_name": getattr(p, "legal_company_name", None),
        "trade_name": getattr(p, "trade_name", None),
        "registered_address": getattr(p, "registered_address", None),
        "pin_code": getattr(p, "pin_code", None),
        "gstin": getattr(p, "gstin", None),
        "pan": getattr(p, "pan", None),
        "authorized_signatory_name": getattr(p, "authorized_signatory_name", None),
        "designation": getattr(p, "designation", None),
        "data_1": getattr(p, "data_1", None),
        "data_2": getattr(p, "data_2", None),
        "data_3": getattr(p, "data_3", None),
        "status": p.status,
        "allow_member_upload": getattr(p, "allow_member_upload", True),
        "card_logo_key": getattr(p, "card_logo_key", None),
        "card_color": getattr(p, "card_color", None),
        "email": str(user.email) if user else None,
        "mobile_no": user.mobile_no if user else None,
        "api_key": p.api_key,
        "api_rate_limit": p.api_rate_limit,
        "created_at": p.created_at.isoformat() if p.created_at else None,
    }


def _enrich(partners: list) -> list:
    uq = UserQuery()
    mq = MemberQuery()
    from ...db.queries.activity_query import NotificationQuery
    nq = NotificationQuery()
    result = []
    for p in partners:
        user = uq.get_user_by_id(str(p.user_id))
        d = _partner_dict(p, user)
        d["member_count"] = mq.count_by_partner(str(p.id))
        d["unread_notification_count"] = nq.unread_count(str(p.user_id))
        result.append(d)
    return result


@admin_partners_router.post("/list", response_model=ResponseModel)
async def list_partners(
    body: PartnerListRequest,
    request: Request,
    _=Depends(_require_superadmin),
):
    """
    Paginated partner list.

    filters[] supports: name, partner_type, status, city
    """
    pq = PartnerQuery()
    total, partners = pq.list_paginated(body)
    return ResponseModel.ok(data={
        "data": _enrich(partners),
        "total": total,
        "skip": body.skip,
        "limit": body.limit,
    })


@admin_partners_router.post("", response_model=ResponseModel, status_code=201)
async def create_partner(body: PartnerCreate, request: Request, _=Depends(_require_superadmin)):
    p = PartnerService().create(body)
    if body.plan_ids:
        from ...services.plan_service import PlanService
        ps = PlanService()
        for plan_id in body.plan_ids:
            try:
                ps.link_partner(plan_id, str(p.id))
            except Exception:
                pass
    user = UserQuery().get_user_by_id(str(p.user_id))
    return ResponseModel.ok(data=_partner_dict(p, user))


# ── Column map for partner bulk upload ───────────────────────────────────────
_PARTNER_COL_MAP = {
    "partner type":                "partner_type",
    "legal company name":          "legal_company_name",
    "trade name":                  "trade_name",
    "trade name/brand name":       "trade_name",
    "registered office address":   "registered_address",
    "city":                        "city",
    "state":                       "state",
    "pin code":                    "pin_code",
    "gstin":                       "gstin",
    "pan":                         "pan",
    "authorized signatory name":   "authorized_signatory_name",
    "designation":                 "designation",
    "mobile number":               "mobile_no",
    "email id":                    "email",
    "data 1":                      "data_1",
    "data 2":                      "data_2",
    "data 3":                      "data_3",
}

_PARTNER_MANDATORY = {
    "legal_company_name", "trade_name", "registered_address", "city",
    "state", "pin_code", "gstin", "pan",
    "authorized_signatory_name", "designation", "mobile_no", "email",
}

_GSTIN_RE = re.compile(r"^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z]{1}[1-9A-Z]{1}Z[0-9A-Z]{1}$")
_PAN_RE   = re.compile(r"^[A-Z]{5}[0-9]{4}[A-Z]{1}$")
_PIN_RE   = re.compile(r"^\d{6}$")
_MOB_RE   = re.compile(r"^\+?[\d\s\-()]{7,15}$")


def _validate_partner_row(row_data: dict) -> list:
    """Return list of error strings for a parsed row dict. Empty list = valid."""
    errors = []
    for field in _PARTNER_MANDATORY:
        if not row_data.get(field):
            label = field.replace("_", " ").title()
            errors.append(f"{label} is required")

    gstin = row_data.get("gstin", "")
    if gstin and not _GSTIN_RE.match(gstin.upper()):
        errors.append("Invalid GSTIN format (expected 15 chars, e.g. 27AAPFU0939F1ZV)")

    pan = row_data.get("pan", "")
    if pan and not _PAN_RE.match(pan.upper()):
        errors.append("Invalid PAN format (expected 10 chars, e.g. AAPFU0939F)")

    pin = row_data.get("pin_code", "")
    if pin and not _PIN_RE.match(str(pin)):
        errors.append("Invalid Pin Code — must be exactly 6 digits")

    mob = row_data.get("mobile_no", "")
    if mob and not _MOB_RE.match(mob):
        errors.append("Invalid Mobile Number format")

    return errors


def _parse_partner_excel(contents: bytes):
    """Parse Excel bytes → (header_row, data_rows, col_idx)."""
    wb = openpyxl.load_workbook(filename=io.BytesIO(contents), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise HTTPException(status_code=422, detail="Excel file is empty")
    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    col_idx = {_PARTNER_COL_MAP[h]: i for i, h in enumerate(header) if h in _PARTNER_COL_MAP}
    missing = _PARTNER_MANDATORY - set(col_idx.keys())
    if missing:
        raise HTTPException(
            status_code=422,
            detail=f"Missing mandatory columns: {', '.join(sorted(missing))}"
        )
    return rows[0], rows[1:], col_idx


def _row_to_dict(row, col_idx: dict) -> dict:
    def cell(field):
        idx = col_idx.get(field)
        if idx is None:
            return None
        v = row[idx] if idx < len(row) else None
        return str(v).strip() if v is not None else None
    return {field: cell(field) for field in list(_PARTNER_COL_MAP.values())}


@admin_partners_router.post("/bulk-upload", response_model=ResponseModel, status_code=201)
async def bulk_upload_partners(
    request: Request,
    file: UploadFile = File(...),
    plan_ids: Optional[str] = Form(None),
    _=Depends(_require_superadmin),
):
    """
    Upload an Excel file to bulk-create partners.
    Mandatory columns: Legal Company Name, Trade Name/Brand Name, Registered Office Address,
    City, State, Pin Code, GSTIN, PAN, Authorized Signatory Name, Designation,
    Mobile Number, Email ID.
    Optional: Partner Type, Data 1, Data 2, Data 3.
    """
    contents = await file.read()
    try:
        header_row, data_rows, col_idx = _parse_partner_excel(contents)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=422, detail="Invalid Excel file — must be .xlsx format")

    svc = PartnerService()
    uq = UserQuery()
    admin_id = getattr(request.state, "user_id", "admin")
    ip = request.client.host if request.client else None
    results = {"created": [], "skipped": [], "errors": []}

    for row_num, row in enumerate(data_rows, start=2):
        rd = _row_to_dict(row, col_idx)
        row_errors = _validate_partner_row(rd)
        if row_errors:
            results["errors"].append({
                "row": row_num,
                "email": rd.get("email"),
                "errors": row_errors,
            })
            continue
        try:
            partner_data = PartnerCreate(
                name=rd["authorized_signatory_name"],
                email=rd["email"],
                mobile_no=rd["mobile_no"],
                partner_type=rd.get("partner_type") or "Broker",
                legal_company_name=rd["legal_company_name"],
                trade_name=rd["trade_name"],
                registered_address=rd["registered_address"],
                city=rd["city"],
                state=rd["state"],
                pin_code=rd["pin_code"],
                gstin=rd["gstin"].upper(),
                pan=rd["pan"].upper(),
                authorized_signatory_name=rd["authorized_signatory_name"],
                designation=rd["designation"],
                data_1=rd.get("data_1"),
                data_2=rd.get("data_2"),
                data_3=rd.get("data_3"),
            )
            p = svc.create(partner_data)
            if plan_ids:
                from ...services.plan_service import PlanService
                ps = PlanService()
                for pid in plan_ids.split(","):
                    pid = pid.strip()
                    if pid:
                        try:
                            ps.link_partner(pid, str(p.id))
                        except Exception:
                            pass
            results["created"].append({"row": row_num, "email": rd["email"], "id": str(p.id)})
        except HTTPException as e:
            results["skipped"].append({"row": row_num, "email": rd.get("email"), "reason": e.detail})
        except Exception as e:
            results["errors"].append({"row": row_num, "email": rd.get("email"), "errors": [str(e)]})

    AuditService().log(
        actor_id=admin_id, actor_type="admin",
        action="bulk_upload",
        entity_type="partner",
        note=f"Partner bulk upload: {len(results['created'])} created, {len(results['skipped'])} skipped, {len(results['errors'])} errors",
        ip_address=ip,
    )
    return ResponseModel.ok(data={"total_rows": len(data_rows), **results})


@admin_partners_router.post("/bulk-upload/report")
async def partner_bulk_upload_report(
    request: Request,
    file: UploadFile = File(...),
    _=Depends(_require_superadmin),
):
    """
    Dry-run the same file and return an annotated Excel with:
    - Green row = valid
    - Red row = has errors
    - 'Status' column appended
    - 'Error Details' column appended
    """
    contents = await file.read()
    try:
        header_row, data_rows, col_idx = _parse_partner_excel(contents)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=422, detail="Invalid Excel file")

    # Build annotated workbook
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Partner Upload Report"

    GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    RED_FILL   = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    HDR_FILL   = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    GREEN_FONT = Font(color="276128", bold=False)
    RED_FONT   = Font(color="9C0006", bold=False)
    HDR_FONT   = Font(bold=True)

    # Write header row + two extra columns
    header_list = list(header_row) + ["Status", "Error Details"]
    for col_idx_h, val in enumerate(header_list, start=1):
        cell = ws_out.cell(row=1, column=col_idx_h, value=val)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT

    status_col = len(header_list) - 1  # 1-indexed
    error_col  = len(header_list)

    # Process each data row
    for row_num, row in enumerate(data_rows, start=2):
        rd = _row_to_dict(row, col_idx)
        row_errors = _validate_partner_row(rd)
        fill   = RED_FILL   if row_errors else GREEN_FILL
        font   = RED_FONT   if row_errors else GREEN_FONT
        status = "Error ✗"  if row_errors else "Success ✓"
        err_msg = "; ".join(row_errors) if row_errors else ""

        for col_i, val in enumerate(row, start=1):
            cell = ws_out.cell(row=row_num, column=col_i, value=val)
            cell.fill = fill
            cell.font = font

        # Pad any missing cells up to status column
        for col_i in range(len(row) + 1, status_col):
            ws_out.cell(row=row_num, column=col_i, value="").fill = fill

        ws_out.cell(row=row_num, column=status_col, value=status).fill = fill
        err_cell = ws_out.cell(row=row_num, column=error_col, value=err_msg)
        err_cell.fill = fill
        err_cell.font = font

    # Auto-size columns (approximate)
    for col in ws_out.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws_out.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb_out.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=partner_upload_report.xlsx"},
    )


@admin_partners_router.get("/bulk-upload/sample")
async def download_partner_sample_excel(_=Depends(_require_superadmin)):
    """Return a sample Excel file with the correct partner upload column headers."""
    from io import BytesIO
    from openpyxl.styles import Alignment

    HEADERS = [
        "Legal Company Name", "Trade Name/Brand Name", "Registered Office Address",
        "City", "State", "Pin Code", "GSTIN", "PAN",
        "Authorized Signatory Name", "Designation", "Mobile Number", "Email ID",
        "Partner Type", "Data 1", "Data 2", "Data 3",
    ]
    SAMPLE_ROW = [
        "Acme Insurance Pvt Ltd", "Acme Insurance", "123 MG Road, Andheri East",
        "Mumbai", "Maharashtra", "400069", "27AAPFU0939F1ZV", "AAPFU0939F",
        "Rahul Sharma", "Director", "9876543210", "rahul@acmecorp.com",
        "Broker", "", "", "",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Partners"

    # Header row — bold, navy fill
    header_fill = PatternFill(start_color="0A2257", end_color="0A2257", fill_type="solid")
    for col_num, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(len(header) + 4, 18)

    # Sample data row
    for col_num, value in enumerate(SAMPLE_ROW, start=1):
        ws.cell(row=2, column=col_num, value=value)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=partner_upload_sample.xlsx"},
    )


@admin_partners_router.get("/{partner_id}", response_model=ResponseModel)
async def get_partner(partner_id: UUID, request: Request, _=Depends(_require_superadmin)):
    p = PartnerService().get_by_id(str(partner_id))
    user = UserQuery().get_user_by_id(str(p.user_id))
    return ResponseModel.ok(data=_partner_dict(p, user))


@admin_partners_router.patch("/{partner_id}", response_model=ResponseModel)
async def update_partner(partner_id: UUID, body: PartnerUpdate, request: Request,
                         _=Depends(_require_superadmin)):
    p = PartnerService().update(str(partner_id), body)
    user = UserQuery().get_user_by_id(str(p.user_id))
    return ResponseModel.ok(data=_partner_dict(p, user))


@admin_partners_router.delete("/{partner_id}", response_model=ResponseModel)
async def delete_partner(partner_id: UUID, request: Request, _=Depends(_require_superadmin)):
    PartnerService().delete(str(partner_id))
    return ResponseModel.ok(data={"message": "Partner deactivated"})


@admin_partners_router.post("/{partner_id}/regenerate-key", response_model=ResponseModel)
async def regen_key(partner_id: UUID, request: Request, _=Depends(_require_superadmin)):
    p = PartnerService().regenerate_api_key(str(partner_id))
    return ResponseModel.ok(data={"api_key": p.api_key})


# ── Membership Card branding (E6/E7/E10, 4th MOM) ────────────────────────────
# Partners get a controlled set of branding fields (logo + accent color) instead
# of free-form HTML — safer for non-technical users, no HTML-to-PDF engine needed.

@admin_partners_router.post("/{partner_id}/card-logo", response_model=ResponseModel, status_code=201)
async def upload_card_logo(partner_id: UUID, request: Request,
                           file: UploadFile = File(...), _=Depends(_require_superadmin)):
    from ...storage import get_storage
    if file.content_type not in ("image/png", "image/jpeg", "image/jpg", "image/webp"):
        raise HTTPException(status_code=422, detail="Logo must be a PNG, JPEG, or WEBP image")
    contents = await file.read()
    if len(contents) > 2 * 1024 * 1024:
        raise HTTPException(status_code=422, detail="Logo must be under 2MB")
    ext = (file.filename or "logo.png").rsplit(".", 1)[-1].lower()
    key = f"partners/{partner_id}/card-logo.{ext}"
    get_storage().upload(key, contents, content_type=file.content_type)
    PartnerQuery().update(str(partner_id), card_logo_key=key)
    return ResponseModel.ok(data={"card_logo_key": key})


@admin_partners_router.get("/{partner_id}/card-logo/view")
async def view_card_logo(partner_id: UUID, _=Depends(_require_superadmin)):
    from fastapi import Response
    from ...storage import get_storage
    p = PartnerService().get_by_id(str(partner_id))
    if not p.card_logo_key:
        raise HTTPException(status_code=404, detail="No logo uploaded for this partner")
    data = get_storage().download(p.card_logo_key)
    ext = p.card_logo_key.rsplit(".", 1)[-1].lower()
    media_type = "image/png" if ext == "png" else "image/webp" if ext == "webp" else "image/jpeg"
    return Response(content=data, media_type=media_type)


@admin_partners_router.get("/{partner_id}/card-preview")
async def preview_membership_card(partner_id: UUID, _=Depends(_require_superadmin)):
    """Download a Membership Card PDF with dummy placeholder values, rendered with
    this partner's current branding (logo + color) — so they can verify it before go-live."""
    from fastapi import Response
    from datetime import date, timedelta
    from ...storage import get_storage
    from ...services.pdf_service import PdfService

    p = PartnerService().get_by_id(str(partner_id))
    logo_bytes = None
    if p.card_logo_key:
        try:
            logo_bytes = get_storage().download(p.card_logo_key)
        except Exception:
            logo_bytes = None

    class _DummyPlan:
        benefit_family = 4
        benefit_slots = 3
        benefit_claim = "Priority"
        benefit_aiqa = True
        benefit_teleconsult_sessions = 2
        benefit_hospital_cash = True
        benefit_wellness_sessions = 1
        benefit_emergency_assist = True

    today = date.today()
    pdf_bytes = PdfService().generate_membership_card_pdf(
        member_name="John Doe",
        member_email="john.doe@example.com",
        partner_name=p.name,
        partner_type=p.partner_type,
        plan_name="Sample Plan",
        plan=_DummyPlan(),
        membership_number="MEM-SAMPLE1",
        start_date=str(today),
        end_date=str(today + timedelta(days=365)),
        partner_address=p.registered_address,
        partner_logo_bytes=logo_bytes,
        partner_color=p.card_color,
    )
    return Response(
        content=pdf_bytes, media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="membership_card_preview.pdf"'},
    )


@admin_partners_router.get("/{partner_id}/plans", response_model=ResponseModel)
async def partner_plans(partner_id: UUID, request: Request, _=Depends(_require_superadmin)):
    plans = PlanService().list_for_partner(str(partner_id))
    return ResponseModel.ok(data=[_plan_to_dict(p) for p in plans])


@admin_partners_router.get("/{partner_id}/plans-overview", response_model=ResponseModel)
async def partner_plans_overview(partner_id: UUID, request: Request, _=Depends(_require_superadmin)):
    """All active plans with linked status and member count for this partner."""
    from ...db.models.plan import MembershipPlan
    from ...db.models.partner import PartnerPlan
    from ...db.models.member import MemberEnrollment
    from ...db.session import session_scope
    from sqlalchemy import func

    with session_scope() as session:
        all_plans = session.query(MembershipPlan).filter(
            MembershipPlan.status == "Active",
            MembershipPlan.is_deleted == False,
        ).order_by(MembershipPlan.name).all()

        linked_ids = {
            str(row.plan_id)
            for row in session.query(PartnerPlan).filter(
                PartnerPlan.partner_id == partner_id
            ).all()
        }

        # member count per plan for this partner in one query
        counts = dict(
            session.query(MemberEnrollment.plan_id, func.count(MemberEnrollment.id))
            .filter(MemberEnrollment.partner_id == partner_id)
            .group_by(MemberEnrollment.plan_id)
            .all()
        )

        result = []
        for p in all_plans:
            result.append({
                "id": str(p.id),
                "name": p.name,
                "plan_type": p.plan_type,
                "status": p.status,
                "price": float(p.price) if p.price else None,
                "cycle": p.cycle,
                "description": getattr(p, "tagline", None) or getattr(p, "info_text", None),
                "linked": str(p.id) in linked_ids,
                "member_count": counts.get(p.id, 0),
            })

    return ResponseModel.ok(data=result)


@admin_partners_router.post("/{partner_id}/members/list", response_model=ResponseModel)
async def list_partner_members(
    partner_id: UUID, body: MemberListRequest, request: Request, _=Depends(_require_superadmin)
):
    """Members enrolled under a specific partner, with their plan and policy count."""
    mq = MemberQuery()
    pq = PolicyQuery()
    from ...db.queries.plan_query import PlanQuery
    plan_q = PlanQuery()
    body.partner_id = str(partner_id)
    total, users = mq.list_members_paginated(body, partner_id=str(partner_id))
    result = []
    for u in users:
        enrollment = mq.get_enrollment(str(u.id), str(partner_id))
        plan_name = None
        if enrollment and enrollment.plan_id:
            plan = plan_q.get_by_id(str(enrollment.plan_id))
            plan_name = plan.name if plan else None
        policies = pq.list_by_user_partner(str(u.id), str(partner_id))
        family = mq.list_family(str(u.id))
        result.append({
            "id": str(u.id),
            "name": u.name,
            "email": u.email,
            "mobile_no": u.mobile_no,
            "is_active": u.is_active,
            "enrollment_status": enrollment.status if enrollment else None,
            "plan_id": str(enrollment.plan_id) if enrollment and enrollment.plan_id else None,
            "plan_name": plan_name,
            "policy_count": len(policies),
            "family_count": len(family),
        })
    return ResponseModel.ok(data={"data": result, "total": total, "skip": body.skip, "limit": body.limit})


@admin_partners_router.get("/{partner_id}/notifications", response_model=ResponseModel)
async def list_partner_notifications(
    partner_id: UUID,
    request: Request,
    skip: int = 0,
    limit: int = 50,
    _=Depends(_require_superadmin),
):
    """Admin: list notifications sent to a specific partner."""
    pq = PartnerQuery()
    partner = pq.get_by_id(str(partner_id))
    if not partner:
        raise HTTPException(status_code=404, detail="Partner not found")
    from ...db.queries.activity_query import NotificationQuery
    nq = NotificationQuery()
    total, notifs = nq.list_for_user(str(partner.user_id), skip=skip, limit=limit)
    unread_count = nq.unread_count(str(partner.user_id))
    return ResponseModel.ok(data={
        "data": [
            {
                "id": str(n.id),
                "type": n.type,
                "title": n.title,
                "body": n.body,
                "is_read": n.is_read,
                "created_at": n.created_at.isoformat() if n.created_at else None,
            }
            for n in notifs
        ],
        "total": total,
        "unread_count": unread_count,
        "skip": skip,
        "limit": limit,
    })


@admin_partners_router.patch("/{partner_id}/notifications/mark-all-read", response_model=ResponseModel)
async def mark_partner_notifications_read(
    partner_id: UUID,
    request: Request,
    _=Depends(_require_superadmin),
):
    """Admin marks all notifications for a partner as read."""
    pq = PartnerQuery()
    partner = pq.get_by_id(str(partner_id))
    if not partner:
        raise HTTPException(status_code=404, detail="Partner not found")
    from ...db.queries.activity_query import NotificationQuery
    nq = NotificationQuery()
    count = nq.mark_all_read(str(partner.user_id))
    return ResponseModel.ok(data={"marked_read": count})


@admin_partners_router.post("/{partner_id}/members", response_model=ResponseModel, status_code=201)
async def admin_add_member_to_partner(partner_id: UUID, request: Request, _=Depends(_require_superadmin)):
    """Admin adds a new member to a specific partner."""
    import json as _json
    from ...schemas.member import MemberCreate
    from ...services.member_service import MemberService
    body_bytes = await request.body()
    data = _json.loads(body_bytes)
    body = MemberCreate(**data)
    body.partner_id = str(partner_id)
    result = MemberService().create_member(body)
    user = result["user"]
    enrollment = result["enrollment"]
    from ...services.notification_helper import notify_partner
    notify_partner(
        partner_id=str(partner_id),
        type="new_member",
        title=f"New member added: {user.name}",
        body=f"A new member ({user.email}) has been enrolled under your account.",
        ref_id=str(user.id),
        ref_type="member",
    )
    return ResponseModel.ok(data={
        "id": str(user.id), "email": user.email, "name": user.name,
        "enrollment": {
            "plan_id": str(enrollment.plan_id),
            "status": enrollment.status,
            "end_date": str(enrollment.end_date),
        },
    })


_MEMBER_COL_MAP = {
    "email": "email", "email id": "email",
    "name": "name", "full name": "name",
    "mobile": "mobile_no", "mobile number": "mobile_no", "mobile no": "mobile_no",
    "gender": "gender",
    "address": "address_line", "address line": "address_line",
    "city": "address_city",
    "state": "address_state",
    "pin": "address_pin", "pin code": "address_pin", "pincode": "address_pin",
    "sale date": "sale_date",
    "sales channel": "sales_channel",
    "branch code": "branch_code",
    "salesperson": "salesperson_name", "salesperson name": "salesperson_name",
    "employee code": "employee_code",
    "data 1": "data1", "data1": "data1",
    "data 2": "data2", "data2": "data2",
    "data 3": "data3", "data3": "data3",
    "plan": "plan_name", "plan name": "plan_name",
}

_MEMBER_MANDATORY = {"email", "name", "mobile_no"}


def _resolve_plan_for_row(partner_id: str, plan_name: Optional[str], default_plan_id: Optional[str]) -> tuple:
    """Resolve the plan to enroll this row into. Row's own 'Plan Name' column wins;
    falls back to the upload's default plan if the row doesn't specify one.
    Returns (plan_id, error_message)."""
    from ...db.queries.plan_query import PlanQuery
    from ...db.models.partner import PartnerPlan
    from ...db.session import session_scope

    if plan_name:
        plan = PlanQuery().get_by_name(plan_name)
        if not plan or plan.status != "Active":
            return None, f"Plan '{plan_name}' not found or not active"
        if plan.plan_type == "partner":
            with session_scope() as s:
                linked = s.query(PartnerPlan).filter(
                    PartnerPlan.partner_id == partner_id, PartnerPlan.plan_id == plan.id,
                ).first()
            if not linked:
                return None, f"Plan '{plan_name}' is not available to this partner"
        return str(plan.id), None
    if default_plan_id:
        return default_plan_id, None
    return None, "Plan is required — add a value in the 'Plan Name' column, or select a default plan before uploading"
_MEMBER_MOB_RE = re.compile(r"^\+?[\d\s\-()]{7,15}$")
_MEMBER_PIN_RE = re.compile(r"^\d{6}$")


def _validate_member_row(rd: dict) -> list:
    errors = []
    for f in _MEMBER_MANDATORY:
        if not rd.get(f):
            errors.append(f"{f.replace('_', ' ').title()} is required")
    mob = rd.get("mobile_no", "")
    if mob and not _MEMBER_MOB_RE.match(mob):
        errors.append("Invalid Mobile Number format")
    pin = rd.get("address_pin", "")
    if pin and not _MEMBER_PIN_RE.match(str(pin)):
        errors.append("Invalid PIN Code — must be 6 digits")
    gender = rd.get("gender", "")
    if gender and gender not in ("Male", "Female", "Other"):
        errors.append("Gender must be Male, Female, or Other")
    return errors


def _parse_member_excel(contents: bytes):
    wb = openpyxl.load_workbook(filename=io.BytesIO(contents), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise HTTPException(status_code=422, detail="Excel file is empty")
    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    col_idx = {_MEMBER_COL_MAP[h]: i for i, h in enumerate(header) if h in _MEMBER_COL_MAP}
    missing = _MEMBER_MANDATORY - set(col_idx.keys())
    if missing:
        raise HTTPException(
            status_code=422,
            detail=f"Missing mandatory columns: {', '.join(sorted(missing))}"
        )
    return rows[0], rows[1:], col_idx


def _member_row_to_dict(row, col_idx: dict) -> dict:
    def cell(field):
        idx = col_idx.get(field)
        if idx is None:
            return None
        v = row[idx] if idx < len(row) else None
        return str(v).strip() if v is not None else None
    return {field: cell(field) for field in list(_MEMBER_COL_MAP.values())}


@admin_partners_router.post("/{partner_id}/members/bulk-upload", response_model=ResponseModel)
async def admin_bulk_upload_members_to_partner(
    partner_id: UUID,
    file: UploadFile = File(...),
    plan_id: Optional[str] = Form(None),
    request: Request = None,
    _=Depends(_require_superadmin),
):
    """Admin bulk uploads members to a specific partner from Excel."""
    from ...services.member_service import MemberService
    from ...schemas.member import MemberCreate

    contents = await file.read()
    try:
        header_row, data_rows, col_idx = _parse_member_excel(contents)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=422, detail="Invalid Excel file — must be .xlsx format")

    svc = MemberService()
    results = {"total_rows": len(data_rows), "created": [], "skipped": [], "errors": []}

    for row_num, row in enumerate(data_rows, start=2):
        rd = _member_row_to_dict(row, col_idx)
        row_errors = _validate_member_row(rd)
        if row_errors:
            results["errors"].append({"row": row_num, "email": rd.get("email"), "errors": row_errors})
            continue
        email = rd["email"]
        resolved_plan_id, plan_error = _resolve_plan_for_row(str(partner_id), rd.get("plan_name"), plan_id)
        if plan_error:
            results["errors"].append({"row": row_num, "email": email, "errors": [plan_error]})
            continue
        try:
            body = MemberCreate(
                email=email,
                name=rd.get("name") or email.split("@")[0],
                mobile_no=rd.get("mobile_no") or "0000000000",
                gender=rd.get("gender") or None,
                address_line=rd.get("address_line") or None,
                address_city=rd.get("address_city") or None,
                address_state=rd.get("address_state") or None,
                address_pin=rd.get("address_pin") or None,
                sale_date=rd.get("sale_date") or None,
                sales_channel=rd.get("sales_channel") or None,
                branch_code=rd.get("branch_code") or None,
                salesperson_name=rd.get("salesperson_name") or None,
                employee_code=rd.get("employee_code") or None,
                data1=rd.get("data1") or None,
                data2=rd.get("data2") or None,
                data3=rd.get("data3") or None,
                partner_id=str(partner_id),
                plan_id=resolved_plan_id,
            )
            svc.create_member(body)
            results["created"].append({"row": row_num, "email": email})
        except HTTPException as e:
            results["skipped"].append({"row": row_num, "email": email, "reason": e.detail})
        except Exception as exc:
            results["errors"].append({"row": row_num, "email": email, "errors": [str(exc)]})

    if results["created"]:
        from ...services.notification_helper import notify_partner
        notify_partner(
            partner_id=str(partner_id),
            type="new_member",
            title=f"{len(results['created'])} new member(s) added via bulk upload",
            body=f"{len(results['created'])} member(s) were enrolled under your account.",
            ref_id=str(partner_id),
            ref_type="partner",
        )
    return ResponseModel.ok(data=results)


@admin_partners_router.get("/{partner_id}/members/bulk-upload/sample")
async def admin_member_bulk_sample(partner_id: UUID, _=Depends(_require_superadmin)):
    """Return a sample Excel for member bulk upload."""
    from openpyxl.styles import Alignment
    from ...db.queries.plan_query import PlanQuery

    available_plans = PlanQuery().list_for_partner(str(partner_id), active_only=True)
    sample_plan_name = available_plans[0].name if available_plans else "Gold Plan"

    HEADERS = [
        "Email ID", "Name", "Mobile Number", "Gender", "Plan Name",
        "Address", "City", "State", "PIN Code",
        "Sale Date", "Sales Channel", "Branch Code", "Salesperson Name", "Employee Code",
        "Data 1", "Data 2", "Data 3",
    ]
    SAMPLE = [
        "john.doe@example.com", "John Doe", "9876543210", "Male", sample_plan_name,
        "123 MG Road", "Mumbai", "Maharashtra", "400001",
        "2024-01-15", "Direct", "BRN001", "Jane Smith", "EMP123",
        "", "", "",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Members"
    hdr_fill = PatternFill(start_color="0A2257", end_color="0A2257", fill_type="solid")
    for col_num, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_num, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(len(h) + 4, 18)
    for col_num, v in enumerate(SAMPLE, start=1):
        ws.cell(row=2, column=col_num, value=v)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=member_upload_sample.xlsx"},
    )


@admin_partners_router.post("/{partner_id}/members/bulk-upload/report")
async def admin_member_bulk_report(
    partner_id: UUID,
    file: UploadFile = File(...),
    _=Depends(_require_superadmin),
):
    """Dry-run member upload and return annotated Excel with error rows highlighted red."""
    contents = await file.read()
    try:
        header_row, data_rows, col_idx = _parse_member_excel(contents)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=422, detail="Invalid Excel file")

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Member Upload Report"

    GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    RED_FILL   = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    HDR_FILL   = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    GREEN_FONT = Font(color="276128")
    RED_FONT   = Font(color="9C0006")
    HDR_FONT   = Font(bold=True)

    header_list = list(header_row) + ["Status", "Error Details"]
    for col_idx_h, val in enumerate(header_list, start=1):
        cell = ws_out.cell(row=1, column=col_idx_h, value=val)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT

    status_col = len(header_list) - 1
    error_col  = len(header_list)

    for row_num, row in enumerate(data_rows, start=2):
        rd = _member_row_to_dict(row, col_idx)
        row_errors = _validate_member_row(rd)
        fill   = RED_FILL   if row_errors else GREEN_FILL
        font   = RED_FONT   if row_errors else GREEN_FONT
        status = "Error ✗"  if row_errors else "OK ✓"
        err_msg = "; ".join(row_errors) if row_errors else ""
        for col_i, val in enumerate(row, start=1):
            cell = ws_out.cell(row=row_num, column=col_i, value=val)
            cell.fill = fill
            cell.font = font
        for col_i in range(len(row) + 1, status_col):
            ws_out.cell(row=row_num, column=col_i, value="").fill = fill
        ws_out.cell(row=row_num, column=status_col, value=status).fill = fill
        err_cell = ws_out.cell(row=row_num, column=error_col, value=err_msg)
        err_cell.fill = fill
        err_cell.font = font

    for col in ws_out.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws_out.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb_out.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=member_upload_report.xlsx"},
    )


@admin_partners_router.post("/{partner_id}/policies/list", response_model=ResponseModel)
async def list_partner_policies(
    partner_id: UUID, body: PolicyListRequest, request: Request, _=Depends(_require_superadmin)
):
    """Policies uploaded under a specific partner, with member info."""
    pq = PolicyQuery()
    uq = UserQuery()
    ptq = PolicyTypeQuery()
    body.partner_id = str(partner_id)
    total, policies = pq.list_paginated_all(body, partner_id=str(partner_id), user_id=body.user_id)
    user_cache: dict = {}
    pt_cache: dict = {}
    result = []
    for p in policies:
        uid = str(p.user_id)
        tid = str(p.policy_type_id)
        if uid not in user_cache:
            user_cache[uid] = uq.get_user_by_id(uid)
        if tid not in pt_cache:
            pt_cache[tid] = ptq.get_by_id(tid)
        u = user_cache[uid]
        pt = pt_cache[tid]
        result.append({
            "id": str(p.id),
            "policy_number": p.policy_number,
            "policy_type": pt.name if pt else None,
            "insurer": p.insurer,
            "sum_insured": p.sum_insured,
            "start_date": str(p.start_date) if p.start_date else None,
            "end_date": str(p.end_date) if p.end_date else None,
            "status": p.status,
            "file_name": p.file_name,
            "member_id": uid,
            "member_name": u.name if u else None,
            "member_email": u.email if u else None,
            "has_file": bool(p.storage_key),
            "created_at": p.created_at.isoformat() if p.created_at else None,
        })
    return ResponseModel.ok(data={"data": result, "total": total, "skip": body.skip, "limit": body.limit})


# ── Partner Change Requests ───────────────────────────────────────────────────

def _cr_dict(cr) -> dict:
    return {
        "id":               str(cr.id),
        "partner_id":       str(cr.partner_id),
        "requested_fields": cr.requested_fields,
        "reason":           cr.reason,
        "status":           cr.status,
        "admin_note":       cr.admin_note,
        "reviewed_by":      cr.reviewed_by,
        "reviewed_at":      cr.reviewed_at.isoformat() if cr.reviewed_at else None,
        "created_at":       cr.created_at.isoformat() if cr.created_at else None,
    }


class CRReviewBody(BaseModel):
    admin_note: Optional[str] = None


@admin_partners_router.get("/{partner_id}/change-requests", response_model=ResponseModel)
async def list_partner_change_requests(
    partner_id: UUID, request: Request,
    status: str = None, skip: int = 0, limit: int = 50,
    _=Depends(_require_superadmin),
):
    pq = PartnerQuery()
    total, rows = pq.list_change_requests(
        partner_id=str(partner_id), status=status, skip=skip, limit=limit
    )
    return ResponseModel.ok(data={
        "data": [_cr_dict(cr) for cr in rows],
        "total": total, "skip": skip, "limit": limit,
    })


@admin_partners_router.post("/change-requests/{cr_id}/approve", response_model=ResponseModel)
async def approve_partner_change_request(
    cr_id: str, body: CRReviewBody, request: Request, _=Depends(_require_superadmin)
):
    admin_id = request.state.user_id if hasattr(request.state, "user_id") else "admin"
    cr = PartnerService().approve_change_request(cr_id, admin_id=admin_id, admin_note=body.admin_note)
    return ResponseModel.ok(data=_cr_dict(cr))


@admin_partners_router.post("/change-requests/{cr_id}/reject", response_model=ResponseModel)
async def reject_partner_change_request(
    cr_id: str, body: CRReviewBody, request: Request, _=Depends(_require_superadmin)
):
    admin_id = request.state.user_id if hasattr(request.state, "user_id") else "admin"
    cr = PartnerService().reject_change_request(cr_id, admin_id=admin_id, admin_note=body.admin_note)
    return ResponseModel.ok(data=_cr_dict(cr))
