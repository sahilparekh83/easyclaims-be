import io
import re
from typing import Optional
from uuid import UUID
from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import Response, StreamingResponse
from pydantic import BaseModel
import openpyxl
from openpyxl.styles import PatternFill, Font
from ...schemas.base import ResponseModel
from ...storage import get_storage
from ...schemas.member import MemberCreate
from ...schemas.list_request import MemberListRequest
from ...services.member_service import MemberService
from ...db.queries.user_query import UserQuery
from ...db.queries.member_query import MemberQuery
from ...db.queries.policy_query import PolicyQuery
from ...db.queries.policy_type_query import PolicyTypeQuery
from ...db.queries.activity_query import ActivityQuery, NotificationQuery
from ...db.queries.plan_query import PlanQuery
from ...db.queries.enrollment_history_query import EnrollmentHistoryQuery
from ...db.queries.user_query import UserQuery
from ..deps import _require_partner
from ...db.queries.activity_query import PolicyFamilyQuery as _PFQ

partner_members_router = APIRouter()


def _family_with_counts(family_members):
    pfq = _PFQ()
    result = []
    for f in family_members:
        count = len(pfq.list_by_family_member(str(f.id)))
        result.append({
            "id": str(f.id), "name": f.name, "relation": f.relation,
            "gender": f.gender, "dob": str(f.dob) if f.dob else None,
            "coverage_type": f.coverage_type, "policy_count": count,
        })
    return result


@partner_members_router.post("/list", response_model=ResponseModel)
async def list_members(
    body: MemberListRequest,
    request: Request,
    partner=Depends(_require_partner),
):
    """
    Paginated list of members enrolled under this partner.

    filters[] supports: name, email, mobile_no

    Response shape:
    {
      "data": [
        {
          "id": "...",
          "name": "...",
          "email": "...",
          "mobile_no": "...",
          "enrollment": { plan_id, status, end_date }
        }
      ],
      "total": 12,
      "skip": 0,
      "limit": 10
    }
    """
    mq = MemberQuery()
    pq_plan = PlanQuery()
    pq_policy = PolicyQuery()
    total, users = mq.list_members_paginated(body, partner_id=str(partner.id))

    aq = ActivityQuery()
    result = []
    for u in users:
        enrollment = mq.get_enrollment(str(u.id), str(partner.id))
        activity = aq.get_activity(str(u.id))
        plan_name = None
        if enrollment:
            plan = pq_plan.get_by_id(str(enrollment.plan_id))
            plan_name = plan.name if plan else None
        policy_count = len(pq_policy.list_by_user_partner(str(u.id), str(partner.id)))
        result.append({
            "id": str(u.id),
            "email": u.email,
            "name": u.name,
            "mobile_no": u.mobile_no,
            "is_active": u.is_active,
            "has_logged_in": activity.has_logged_in if activity else False,
            "last_login_at": activity.last_login_at.isoformat() if activity and activity.last_login_at else None,
            "policy_count": policy_count,
            "enrollment": {
                "plan_id": str(enrollment.plan_id),
                "plan_name": plan_name,
                "status": enrollment.status,
                "start_date": str(enrollment.start_date) if enrollment.start_date else None,
                "end_date": str(enrollment.end_date) if enrollment.end_date else None,
            } if enrollment else None,
        })

    return ResponseModel.ok(data={
        "data": result,
        "total": total,
        "skip": body.skip,
        "limit": body.limit,
    })


@partner_members_router.post("", response_model=ResponseModel, status_code=201)
async def create_member(body: MemberCreate, request: Request, partner=Depends(_require_partner)):
    body.partner_id = str(partner.id)
    result = MemberService().create_member(body)
    user = result["user"]
    enrollment = result["enrollment"]
    return ResponseModel.ok(data={
        "id": str(user.id), "email": user.email, "name": user.name,
        "enrollment": {
            "plan_id": str(enrollment.plan_id),
            "status": enrollment.status,
            "end_date": str(enrollment.end_date),
        },
    })


@partner_members_router.get("/{member_id}", response_model=ResponseModel)
async def get_member(member_id: UUID, request: Request, partner=Depends(_require_partner)):
    mq = MemberQuery()
    pq = PolicyQuery()
    ptq = PolicyTypeQuery()
    aq = ActivityQuery()
    planq = PlanQuery()

    enrollment = mq.get_enrollment(str(member_id), str(partner.id))
    if not enrollment:
        raise HTTPException(status_code=404, detail="Member not found under this partner")

    plan = planq.get_by_id(str(enrollment.plan_id)) if enrollment else None
    user = UserQuery().get_user_by_id(str(member_id))
    family = mq.list_family(str(member_id))
    policies = pq.list_by_user_partner(str(member_id), str(partner.id))
    activity = aq.get_activity(str(member_id))

    # Enrich policies with type name, family links, and file info
    from ...db.queries.activity_query import PolicyFamilyQuery
    pfq = PolicyFamilyQuery()
    policy_list = []
    for p in policies:
        pt = ptq.get_by_id(str(p.policy_type_id))
        # Linked family members for this policy
        links = pfq.list_by_policy(str(p.id))
        linked_family = []
        for link in links:
            fm = mq.get_family_member(str(link.family_member_id))
            if fm:
                linked_family.append({"id": str(fm.id), "name": fm.name, "relation": fm.relation})
        policy_list.append({
            "id": str(p.id),
            "policy_number": p.policy_number,
            "policy_type": pt.name if pt else None,
            "insurer": p.insurer,
            "sum_insured": p.sum_insured,
            "start_date": str(p.start_date) if p.start_date else None,
            "end_date": str(p.end_date) if p.end_date else None,
            "status": p.status,
            "file_name": p.file_name,
            "has_file": bool(p.storage_key),
            "linked_family_members": linked_family,
        })

    return ResponseModel.ok(data={
        "id": str(user.id),
        "email": user.email,
        "name": user.name,
        "mobile_no": user.mobile_no,
        "has_logged_in": activity.has_logged_in if activity else False,
        "first_login_at": activity.first_login_at.isoformat() if activity and activity.first_login_at else None,
        "last_login_at": activity.last_login_at.isoformat() if activity and activity.last_login_at else None,
        "login_count": activity.login_count if activity else 0,
        "enrollment": {
            "plan_id": str(enrollment.plan_id),
            "plan_name": plan.name if plan else None,
            "plan_type": plan.plan_type if plan else None,
            "status": enrollment.status,
            "start_date": str(enrollment.start_date) if enrollment.start_date else None,
            "end_date": str(enrollment.end_date) if enrollment.end_date else None,
        },
        "profile": {
            "gender": profile.gender if (profile := mq.get_profile(str(member_id))) else None,
            "dob": str(profile.dob) if profile and profile.dob else None,
            "address_line": profile.address_line if profile else None,
            "address_city": profile.address_city if profile else None,
            "address_state": profile.address_state if profile else None,
            "address_pin": profile.address_pin if profile else None,
            "sale_date": str(profile.sale_date) if profile and profile.sale_date else None,
            "sales_channel": profile.sales_channel if profile else None,
            "branch_code": profile.branch_code if profile else None,
            "salesperson_name": profile.salesperson_name if profile else None,
            "employee_code": profile.employee_code if profile else None,
            "data1": profile.data1 if profile else None,
            "data2": profile.data2 if profile else None,
            "data3": profile.data3 if profile else None,
        },
        "family": _family_with_counts(family),
        "policies": policy_list,
    })


@partner_members_router.get("/{member_id}/policies/{policy_id}/view")
async def view_member_policy_pdf(
    member_id: UUID, policy_id: UUID, request: Request, partner=Depends(_require_partner)
):
    """Partner views a PDF of one of their member's policies."""
    pq = PolicyQuery()
    policy = pq.get_by_id(str(policy_id))
    if not policy or str(policy.partner_id) != str(partner.id) or str(policy.user_id) != str(member_id):
        raise HTTPException(status_code=404, detail="Policy not found")
    if not policy.storage_key:
        raise HTTPException(status_code=404, detail="No file attached to this policy")
    data = get_storage().download(policy.storage_key)
    fname = policy.file_name or "policy.pdf"
    return Response(content=data, media_type="application/pdf",
                    headers={"Content-Disposition": f'inline; filename="{fname}"'})


@partner_members_router.get("/{member_id}/policies/{policy_id}/download")
async def download_member_policy_pdf(
    member_id: UUID, policy_id: UUID, request: Request, partner=Depends(_require_partner)
):
    pq = PolicyQuery()
    policy = pq.get_by_id(str(policy_id))
    if not policy or str(policy.partner_id) != str(partner.id) or str(policy.user_id) != str(member_id):
        raise HTTPException(status_code=404, detail="Policy not found")
    if not policy.storage_key:
        raise HTTPException(status_code=404, detail="No file attached to this policy")
    data = get_storage().download(policy.storage_key)
    fname = policy.file_name or "policy.pdf"
    return Response(content=data, media_type="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@partner_members_router.get("/{member_id}/enrollment/history", response_model=ResponseModel)
async def get_member_enrollment_history(member_id: UUID, request: Request, partner=Depends(_require_partner)):
    """Partner views the enrollment plan change history for one of their members."""
    mq = MemberQuery()
    pq = PlanQuery()
    # Verify member is enrolled under this partner
    enrollment = mq.get_enrollment(str(member_id), str(partner.id))
    if not enrollment:
        raise HTTPException(status_code=404, detail="Member not found under this partner")
    history = EnrollmentHistoryQuery().list_for_enrollment(str(enrollment.id))
    result = []
    for h in history:
        to_plan = pq.get_by_id(str(h.to_plan_id)) if h.to_plan_id else None
        from_plan = pq.get_by_id(str(h.from_plan_id)) if h.from_plan_id else None
        result.append({
            "id": str(h.id),
            "action": h.action,
            "changed_by": h.changed_by,
            "from_plan_name": from_plan.name if from_plan else None,
            "to_plan_name": to_plan.name if to_plan else None,
            "note": h.note,
            "changed_at": h.changed_at.isoformat() if h.changed_at else None,
        })
    return ResponseModel.ok(data=result)


_MEMBER_COL_MAP = {
    "email": "email", "email id": "email",
    "name": "name", "full name": "name",
    "mobile": "mobile_no", "mobile number": "mobile_no", "mobile no": "mobile_no",
    "gender": "gender",
    "address": "address_line", "address line": "address_line",
    "city": "address_city",
    "state": "address_state",
    "pin": "address_pin", "pin code": "address_pin", "pincode": "address_pin",
    "sale date": "sale_date",
    "sales channel": "sales_channel",
    "branch code": "branch_code",
    "salesperson": "salesperson_name", "salesperson name": "salesperson_name",
    "employee code": "employee_code",
    "data 1": "data1", "data1": "data1",
    "data 2": "data2", "data2": "data2",
    "data 3": "data3", "data3": "data3",
}

_MEMBER_MANDATORY = {"email", "name", "mobile_no"}
_MEMBER_MOB_RE = re.compile(r"^\+?[\d\s\-()]{7,15}$")
_MEMBER_PIN_RE = re.compile(r"^\d{6}$")


def _validate_member_row_partner(rd: dict) -> list:
    errors = []
    for f in _MEMBER_MANDATORY:
        if not rd.get(f):
            errors.append(f"{f.replace('_', ' ').title()} is required")
    mob = rd.get("mobile_no", "")
    if mob and not _MEMBER_MOB_RE.match(mob):
        errors.append("Invalid Mobile Number format")
    pin = rd.get("address_pin", "")
    if pin and not _MEMBER_PIN_RE.match(str(pin)):
        errors.append("Invalid PIN Code — must be 6 digits")
    gender = rd.get("gender", "")
    if gender and gender not in ("Male", "Female", "Other"):
        errors.append("Gender must be Male, Female, or Other")
    return errors


def _parse_member_excel_partner(contents: bytes):
    wb = openpyxl.load_workbook(filename=io.BytesIO(contents), read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise HTTPException(status_code=422, detail="Excel file is empty")
    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    col_idx = {_MEMBER_COL_MAP[h]: i for i, h in enumerate(header) if h in _MEMBER_COL_MAP}
    missing = _MEMBER_MANDATORY - set(col_idx.keys())
    if missing:
        raise HTTPException(
            status_code=422,
            detail=f"Missing mandatory columns: {', '.join(sorted(missing))}"
        )
    return rows[0], rows[1:], col_idx


def _member_row_to_dict_partner(row, col_idx: dict) -> dict:
    def cell(field):
        idx = col_idx.get(field)
        if idx is None:
            return None
        v = row[idx] if idx < len(row) else None
        return str(v).strip() if v is not None else None
    return {field: cell(field) for field in list(_MEMBER_COL_MAP.values())}


@partner_members_router.post("/bulk-upload", response_model=ResponseModel)
async def bulk_upload_members(
    file: UploadFile = File(...),
    plan_id: str = Form(...),
    request: Request = None,
    partner=Depends(_require_partner),
):
    """Partner bulk uploads members from Excel. plan_id must be one of the partner's linked plans."""
    from ...db.models.partner import PartnerPlan
    from ...db.session import session_scope

    if plan_id:
        with session_scope() as s:
            linked = s.query(PartnerPlan).filter(
                PartnerPlan.partner_id == partner.id,
                PartnerPlan.plan_id == plan_id,
            ).first()
            if not linked:
                raise HTTPException(status_code=400, detail="Plan not linked to this partner")

    contents = await file.read()
    try:
        header_row, data_rows, col_idx = _parse_member_excel_partner(contents)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=422, detail="Invalid Excel file — must be .xlsx format")

    svc = MemberService()
    results = {"total_rows": len(data_rows), "created": [], "skipped": [], "errors": []}

    for row_num, row in enumerate(data_rows, start=2):
        rd = _member_row_to_dict_partner(row, col_idx)
        row_errors = _validate_member_row_partner(rd)
        if row_errors:
            results["errors"].append({"row": row_num, "email": rd.get("email"), "errors": row_errors})
            continue
        email = rd["email"]
        try:
            body = MemberCreate(
                email=email,
                name=rd.get("name") or email.split("@")[0],
                mobile_no=rd.get("mobile_no") or "0000000000",
                gender=rd.get("gender") or None,
                address_line=rd.get("address_line") or None,
                address_city=rd.get("address_city") or None,
                address_state=rd.get("address_state") or None,
                address_pin=rd.get("address_pin") or None,
                sale_date=rd.get("sale_date") or None,
                sales_channel=rd.get("sales_channel") or None,
                branch_code=rd.get("branch_code") or None,
                salesperson_name=rd.get("salesperson_name") or None,
                employee_code=rd.get("employee_code") or None,
                data1=rd.get("data1") or None,
                data2=rd.get("data2") or None,
                data3=rd.get("data3") or None,
                partner_id=str(partner.id),
                plan_id=plan_id or None,
            )
            svc.create_member(body)
            results["created"].append({"row": row_num, "email": email})
        except HTTPException as e:
            results["skipped"].append({"row": row_num, "email": email, "reason": e.detail})
        except Exception as exc:
            results["errors"].append({"row": row_num, "email": email, "errors": [str(exc)]})

    return ResponseModel.ok(data=results)


@partner_members_router.get("/bulk-upload/sample")
async def partner_member_bulk_sample(partner=Depends(_require_partner)):
    """Return a styled sample Excel for partner member bulk upload."""
    from openpyxl.styles import Alignment

    HEADERS = [
        "Email ID", "Name", "Mobile Number", "Gender",
        "Address", "City", "State", "PIN Code",
        "Sale Date", "Sales Channel", "Branch Code", "Salesperson Name", "Employee Code",
        "Data 1", "Data 2", "Data 3",
    ]
    SAMPLE = [
        "john.doe@example.com", "John Doe", "9876543210", "Male",
        "123 MG Road", "Mumbai", "Maharashtra", "400001",
        "2024-01-15", "Direct", "BRN001", "Jane Smith", "EMP123",
        "", "", "",
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Members"
    hdr_fill = PatternFill(start_color="0A2257", end_color="0A2257", fill_type="solid")
    for col_num, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_num, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(len(h) + 4, 18)
    for col_num, v in enumerate(SAMPLE, start=1):
        ws.cell(row=2, column=col_num, value=v)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=member_upload_sample.xlsx"},
    )


@partner_members_router.post("/bulk-upload/report")
async def partner_member_bulk_report(
    file: UploadFile = File(...),
    partner=Depends(_require_partner),
):
    """Dry-run member upload and return annotated Excel with error rows highlighted red."""
    contents = await file.read()
    try:
        header_row, data_rows, col_idx = _parse_member_excel_partner(contents)
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=422, detail="Invalid Excel file")

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Member Upload Report"

    GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    RED_FILL   = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    HDR_FILL   = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    GREEN_FONT = Font(color="276128")
    RED_FONT   = Font(color="9C0006")
    HDR_FONT   = Font(bold=True)

    header_list = list(header_row) + ["Status", "Error Details"]
    for col_idx_h, val in enumerate(header_list, start=1):
        cell = ws_out.cell(row=1, column=col_idx_h, value=val)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT

    status_col = len(header_list) - 1
    error_col  = len(header_list)

    for row_num, row in enumerate(data_rows, start=2):
        rd = _member_row_to_dict_partner(row, col_idx)
        row_errors = _validate_member_row_partner(rd)
        fill   = RED_FILL   if row_errors else GREEN_FILL
        font   = RED_FONT   if row_errors else GREEN_FONT
        status = "Error ✗"  if row_errors else "OK ✓"
        err_msg = "; ".join(row_errors) if row_errors else ""
        for col_i, val in enumerate(row, start=1):
            cell = ws_out.cell(row=row_num, column=col_i, value=val)
            cell.fill = fill
            cell.font = font
        for col_i in range(len(row) + 1, status_col):
            ws_out.cell(row=row_num, column=col_i, value="").fill = fill
        ws_out.cell(row=row_num, column=status_col, value=status).fill = fill
        err_cell = ws_out.cell(row=row_num, column=error_col, value=err_msg)
        err_cell.fill = fill
        err_cell.font = font

    for col in ws_out.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws_out.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb_out.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=member_upload_report.xlsx"},
    )


